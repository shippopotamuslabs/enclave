#!/usr/bin/env python3
"""Generate X integration roadmap Word document for Enclave."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "X-Pre-Retrieval-And-OAuth-Roadmap.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def main() -> None:
    doc = Document()

    title = doc.add_heading("Enclave X Integration Roadmap", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.add_run("Document version: 1.0\n").italic = True
    meta.add_run("Product: Enclave (Pocket Crew)\n").italic = True
    meta.add_run("Scope: Two planned capabilities — public X pre-retrieval and OAuth user-context tools\n").italic = True

    doc.add_paragraph()
    add_para(
        doc,
        "This document describes two related but distinct X (Twitter) capabilities for Enclave. "
        "They use different authentication models and should not be conflated.",
    )

    # --- Feature 1 ---
    add_heading(doc, "Feature 1: X Host Pre-Retrieval (Pre-Fetch)", 1)

    add_heading(doc, "Summary", 2)
    add_para(
        doc,
        "When the search-intent classifier (and configured skills) indicate that a turn needs live "
        "information, Enclave already runs web pre-retrieval (Tavily/Exa) before the model generates "
        "a reply and injects results under # RELEVANT WEB SEARCH RESULTS: in the system prompt. "
        "X host pre-retrieval extends that pattern to X: run search_x_recent once on the host, shape "
        "the user’s conversational prompt into an X API–compatible keyword query, and inject normalized "
        "results before inference — without requiring the user to sign in with their personal X account.",
    )

    add_heading(doc, "Authentication model", 2)
    add_bullets(
        doc,
        [
            "Uses the existing X Developer API Bearer token stored in Settings (“X API Key”).",
            "Same app-only credentials as the search_x_recent tool executor today.",
            "No OAuth 2.0 user context (PKCE) required for this feature.",
            "Searches public recent posts via X API v2 /2/tweets/search/recent.",
        ],
    )

    add_heading(doc, "Current state (shipped)", 2)
    add_bullets(
        doc,
        [
            "search_x_recent tool: model-invoked via structured XSearchParams and BuildXQueryUseCase.",
            "Intent-gated activation: when search intent fires and X is configured (toggle + API key), "
            "search_x_recent is added to activeTools without disclose_tool_schemas.",
            "Web pre-retrieval: classifier + SearchIntentPolicy drive host web_search before sendPrompt.",
            "Sources pill: web pre-retrieval results surface in chat after persistence/projection fixes.",
        ],
    )

    add_heading(doc, "Problem / why deferred", 2)
    add_para(
        doc,
        "X recent search is not semantic search. Sending the user’s full conversational prompt verbatim "
        "to the API often returns poor or empty results. Exa and Tavily accept natural-language queries; "
        "X expects short, keyword-shaped queries with operators (AND/OR, phrases, @mentions, hashtags). "
        "Host pre-search therefore needs prompt → XSearchParams derivation before calling the API.",
    )

    add_heading(doc, "Planned ticket", 2)
    add_para(doc, "FEAT-x-pre-retrieval-query-shaping", bold=True)

    add_heading(doc, "Technical approach", 2)
    add_bullets(
        doc,
        [
            "DeriveXSearchParamsFromPromptUseCase: map user prompt → XSearchParams (or skip with reason).",
            "Heuristics (planned): max ~6 keywords; max composed query length ~128; skip when prompt is "
            "long conversational text with no quoted phrase, hashtag, or @mention; prefer quoted strings "
            "and proper-noun tokens.",
            "Reuse BuildXQueryUseCase to compose the final query string from structured params.",
            "ChatInferenceRequestPreparer: when ToolIntentActivationPolicy enables X and pre-retrieval "
            "for X, execute search once before sendPrompt (parallel to web pre-retrieval).",
            "PreparedChatInferenceRequest.preRetrievalXResultJson: raw normalized JSON from X search.",
            "Prompt injection: new section # RELEVANT X SEARCH RESULTS: (mirror web block).",
            "Optional: merge X hits into WebSearchSource / Sources UI (same patterns as web pre-fetch).",
        ],
    )

    add_heading(doc, "Layer changes (high level)", 2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Layer"
    hdr[1].text = "Changes"
    rows = [
        (
            ":core:domain",
            "DeriveXSearchParamsFromPromptUseCase; extend ToolIntentActivationPolicy / preparer; "
            "preRetrievalXResultJson on PreparedChatInferenceRequest; prompt composer hook.",
        ),
        (
            ":core:data",
            "Reuse XSearchRepository / XSearchToolExecutorImpl; optional early persist of X sources.",
        ),
        (
            ":feature:chat",
            "Sources pill / projection if X pre-results should appear like web sources.",
        ),
        (
            ":feature:inference",
            "No change to provider-native x_search on Grok 4.20+ (separate xAI-hosted path).",
        ),
    ]
    for layer, change in rows:
        row = table.add_row().cells
        row[0].text = layer
        row[1].text = change

    add_heading(doc, "Non-goals", 2)
    add_bullets(
        doc,
        [
            "User OAuth / personalized timelines / posting as the user.",
            "Changing search-intent classifier weights or threshold in this ticket.",
            "Replacing model-driven search_x_recent calls — pre-fetch is additive.",
        ],
    )

  # --- Feature 2 ---
    doc.add_page_break()
    add_heading(doc, "Feature 2: X OAuth User Context (“Sign in with X”)", 1)

    add_heading(doc, "Summary", 2)
    add_para(
        doc,
        "Some X capabilities require acting as or reading data for a specific signed-in user. "
        "App-only Bearer tokens from the Developer Portal cannot perform write operations or access "
        "private user feeds. Enclave deferred the tools below until OAuth 2.0 User Context with PKCE "
        "is implemented end-to-end.",
    )

    add_heading(doc, "Deferred tools requiring X sign-in", 2)

    add_heading(doc, "x_post_action (write operations)", 3)
    add_bullets(
        doc,
        [
            "Posting tweets, replying to threads, retweeting, quoting, and liking/unliking.",
            "Why PKCE is required: app-only Bearer tokens do not have write permissions; mutating "
            "actions must be authorized by a signed-in user context.",
        ],
    )

    add_heading(doc, "x_notifications (notifications and activity)", 3)
    add_bullets(
        doc,
        [
            "Retrieving the user’s notifications, mentions, and activity feed.",
            "Why PKCE is required: user-specific endpoints (e.g. /2/users/me, mentions, likes timeline) "
            "are restricted to the authenticated user.",
        ],
    )

    add_heading(doc, "Personalized / tailored trends", 3)
    add_bullets(
        doc,
        [
            "Trends customized to the authenticated user’s location and interests.",
            "Example API: /2/users/personalized_trends.",
            "Requires the same user-context OAuth session as the tools above.",
        ],
    )

    add_heading(doc, "A. Developer Portal setup", 2)
    add_bullets(
        doc,
        [
            "Update app registration in the X Developer Portal to enable OAuth 2.0 and PKCE.",
            "Configure a custom redirect URI (e.g. enclave://oauth/x/callback).",
            "Request scopes: tweet.read, tweet.write, users.read, like.write, and offline.access "
            "(required for refresh tokens).",
        ],
    )

    add_heading(doc, "B. Android app authentication flow", 2)
    add_bullets(
        doc,
        [
            "Browser handshake: Custom Tabs or AppAuth to prompt the user to authorize the app.",
            "Callback interceptor: deep-link handling in navigation to capture the authorization code.",
            "Token exchange and storage: exchange code for access_token and refresh_token; store "
            "securely (e.g. ApiKeyManager or dedicated OAuth token store with encryption).",
            "Auto-refresh: X OAuth 2.0 access tokens expire about every 2 hours; HTTP client "
            "interceptor or repository layer refreshes using refresh_token before API calls.",
        ],
    )

    add_heading(doc, "C. Settings UI", 2)
    add_bullets(
        doc,
        [
            "“Connect your X Account” section (distinct from “X API Key” for public search).",
            "Show connected @handle when linked.",
            "“Disconnect / Sign Out” to purge access and refresh tokens.",
        ],
    )

    add_heading(doc, "D. Tool registration and UX safeguards", 2)
    add_bullets(
        doc,
        [
            "Conditional tool inclusion: ChatInferenceRequestPreparer exposes x_post_action and "
            "x_notifications only when a valid user-context token exists — or lists them with a "
            "disabled description so the model can prompt the user to connect X.",
            "Write confirmation gate (critical): x_post_action must never run silently. Tool calls "
            "enter a review state in chat; the API executes only after the user taps Confirm.",
        ],
    )

    add_heading(doc, "Layer changes (high level)", 2)
    table2 = doc.add_table(rows=1, cols=2)
    table2.style = "Table Grid"
    hdr2 = table2.rows[0].cells
    hdr2[0].text = "Layer"
    hdr2[1].text = "Changes"
    oauth_rows = [
        (":core:domain", "XUserSessionPort; OAuth use cases; new ToolDefinitions; gating in preparer."),
        (":core:data", "OAuth PKCE client; token store; user-context X API adapters per tool."),
        (":feature:settings", "Connect / disconnect UI; session state in ViewModel."),
        (":app", "Deep link intent filter for enclave://oauth/x/callback."),
        (":feature:chat", "Write-confirmation composable / ViewModel flow before executor runs."),
    ]
    for layer, change in oauth_rows:
        row = table2.add_row().cells
        row[0].text = layer
        row[1].text = change

    add_heading(doc, "Comparison: Feature 1 vs Feature 2", 2)
    cmp = doc.add_table(rows=1, cols=3)
    cmp.style = "Table Grid"
    ch = cmp.rows[0].cells
    ch[0].text = "Aspect"
    ch[1].text = "X pre-retrieval"
    ch[2].text = "OAuth user tools"
    comparisons = [
        ("Auth", "Developer API Bearer key", "OAuth 2.0 PKCE user tokens"),
        ("Data", "Public recent search", "Private feeds, writes, personalized trends"),
        ("User sign-in", "No", "Yes — Connect X Account"),
        ("Risk", "Query quality / API cost", "Writes, consent, token security"),
        ("Ticket", "FEAT-x-pre-retrieval-query-shaping", "New (OAuth + tools)"),
    ]
    for a, b, c in comparisons:
        r = cmp.add_row().cells
        r[0].text = a
        r[1].text = b
        r[2].text = c

    doc.add_paragraph()
    add_para(
        doc,
        "Recommended implementation order: ship X pre-retrieval first (reuses existing search "
        "infrastructure and API key). OAuth user-context tools are a larger product and security "
        "effort and should follow once write-confirmation UX and token lifecycle are designed.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
